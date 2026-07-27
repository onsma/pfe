import sys, io
sys.path.insert(0, 'd:/BFPME/tmp_pip')
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import openpyxl
wb = openpyxl.load_workbook('BTS_INFO.xlsx')

sf = {}
for row in wb['RU_SousFacteur'].iter_rows(values_only=True):
    if row[0] and row[0] != 'Code':
        sf[row[0]] = row[1]

fac = {}
for row in wb['RU_Facteur'].iter_rows(values_only=True):
    if row[0] and row[0] != 'Code':
        fac[row[0]] = row[1]

sfpm = list(wb['RU_SousFacteurParModele'].iter_rows(values_only=True))

for modele in [1, 2]:
    print('')
    print('========================================')
    print('  MODELE ' + str(modele))
    print('========================================')
    current_fac = None
    rows = [r for r in sfpm if r[2] == modele and r[1] != 'CodeSousFacteur']
    rows_sorted = sorted(rows, key=lambda r: (r[3], r[5]))
    for r in rows_sorted:
        code_sf  = r[1]
        code_mod = r[2]
        code_fac = r[3]
        min_kf   = r[4]
        rang     = r[5]
        prio     = r[6]
        elim     = r[7]
        acc_nd   = r[9]
        val_obl  = r[10]
        if code_fac != current_fac:
            nom_fac = fac.get(code_fac, '?')
            print('')
            print('  --- Facteur ' + str(code_fac) + ': ' + nom_fac + ' ---')
            print('  ' + 'Rang'.ljust(4) + ' ' + 'CodeSF'.ljust(7) + ' ' + 'Sous-Facteur'.ljust(62) + ' ' + 'Min'.ljust(7) + ' ' + 'Prio'.ljust(5) + ' ' + 'Elim'.ljust(5) + ' ND  Obl')
            print('  ' + '-'*4 + ' ' + '-'*7 + ' ' + '-'*62 + ' ' + '-'*7 + ' ' + '-'*5 + ' ' + '-'*5 + ' ' + '-'*4 + ' ' + '-'*4)
            current_fac = code_fac
        libelle = sf.get(code_sf, '?')
        libelle = (libelle[:60] + '..') if len(libelle) > 62 else libelle
        print('  ' + str(rang).ljust(4) + ' ' + str(code_sf).ljust(7) + ' ' + libelle.ljust(62) + ' ' + str(min_kf).ljust(7) + ' ' + str(prio).ljust(5) + ' ' + str(elim).ljust(5) + ' ' + str(acc_nd).ljust(4) + ' ' + str(val_obl))
