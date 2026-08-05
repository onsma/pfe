import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import openpyxl
wb = openpyxl.load_workbook('BTS_INFO.xlsx')

sf = {}
for row in wb['RU_SousFacteur'].iter_rows(values_only=True):
    if row[0] and row[0] != 'Code':
        sf[row[0]] = row[1]

criteres_raw = []
for row in wb['RU_CritereNotes'].iter_rows(values_only=True):
    if row[0] and row[0] != 'Code':
        criteres_raw.append(row)

# Count rules per (modele, sous-facteur)
from collections import defaultdict
rules_per_sf_modele = defaultdict(int)
for r in criteres_raw:
    key = (r[2], r[3])  # (CodeModeleNotation, CodeSousFacteur)
    rules_per_sf_modele[key] += 1

# Count per modele
per_modele = defaultdict(int)
for (mod, sf_code), count in rules_per_sf_modele.items():
    per_modele[mod] += count

print('=== TOTAL REGLES PAR MODELE ===')
for mod, total in sorted(per_modele.items()):
    print(f'  Modele {mod} : {total} regles')
print(f'  TOTAL     : {len(criteres_raw)} regles')

print()
print('=== REGLES PAR SOUS-FACTEUR (Modele 1) ===')
print(f'  {"CodeSF":<8} {"NbRegles":<10} Sous-Facteur')
print(f'  {"-"*8} {"-"*10} {"-"*50}')
m1_rows = [(sf_code, count) for (mod, sf_code), count in rules_per_sf_modele.items() if mod == 1]
for sf_code, count in sorted(m1_rows, key=lambda x: -x[1]):
    libelle = sf.get(sf_code, '?')[:55]
    print(f'  {str(sf_code):<8} {str(count):<10} {libelle}')

print()
print('=== REGLES PAR SOUS-FACTEUR (Modele 2) ===')
print(f'  {"CodeSF":<8} {"NbRegles":<10} Sous-Facteur')
print(f'  {"-"*8} {"-"*10} {"-"*50}')
m2_rows = [(sf_code, count) for (mod, sf_code), count in rules_per_sf_modele.items() if mod == 2]
for sf_code, count in sorted(m2_rows, key=lambda x: -x[1]):
    libelle = sf.get(sf_code, '?')[:55]
    print(f'  {str(sf_code):<8} {str(count):<10} {libelle}')

# Stats
all_counts = list(rules_per_sf_modele.values())
print()
print('=== STATISTIQUES ===')
print(f'  Nb combinaisons (modele x sous-facteur) : {len(rules_per_sf_modele)}')
print(f'  Moyenne regles par sous-facteur          : {sum(all_counts)/len(all_counts):.1f}')
print(f'  Min regles pour un sous-facteur          : {min(all_counts)}')
print(f'  Max regles pour un sous-facteur          : {max(all_counts)}')
